from __future__ import annotations

import html
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


class DocumentGenerationError(RuntimeError):
    pass


class LegalDocumentGenerator:
    def generate_claim_docx(self, text: str, output_path: Path) -> None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        normal_style = document.styles["Normal"]
        normal_style.font.name = "Times New Roman"
        normal_style.font.size = Pt(12)

        for block in self._paragraphs(text):
            paragraph = document.add_paragraph()
            if block.strip().upper() == "ПРЕТЕНЗИЯ":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(block.strip())
                run.bold = True
                run.font.size = Pt(14)
                continue

            lines = block.splitlines()
            for index, line in enumerate(lines):
                paragraph.add_run(line)
                if index < len(lines) - 1:
                    paragraph.add_run().add_break()

        document.save(output_path)

    def generate_claim_pdf(self, text: str, output_path: Path) -> None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        font_name = self._register_cyrillic_font()

        styles = getSampleStyleSheet()
        normal_style = ParagraphStyle(
            "ClaimNormal",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=11,
            leading=15,
            firstLineIndent=8 * mm,
            spaceAfter=6,
        )
        title_style = ParagraphStyle(
            "ClaimTitle",
            parent=normal_style,
            alignment=1,
            fontSize=14,
            leading=18,
            firstLineIndent=0,
            spaceAfter=12,
        )

        story = []
        for block in self._paragraphs(text):
            escaped = "<br/>".join(html.escape(line) for line in block.splitlines())
            style = title_style if block.strip().upper() == "ПРЕТЕНЗИЯ" else normal_style
            story.append(Paragraph(escaped, style))
            story.append(Spacer(1, 4))

        document = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=25 * mm,
            rightMargin=20 * mm,
            topMargin=20 * mm,
            bottomMargin=20 * mm,
        )
        document.build(story)

    @staticmethod
    def _paragraphs(text: str) -> list[str]:
        blocks = [block.strip() for block in text.replace("\r\n", "\n").split("\n\n")]
        return [block for block in blocks if block]

    @staticmethod
    def _register_cyrillic_font() -> str:
        font_name = "ProjectCyrillic"
        if font_name in pdfmetrics.getRegisteredFontNames():
            return font_name

        # ReportLab не встраивает кириллицу стандартными PDF-шрифтами, поэтому ищем системный TTF.
        candidates = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
            "/Library/Fonts/Arial Unicode.ttf",
            "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
            "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
            "C:/Windows/Fonts/arial.ttf",
            "C:/Windows/Fonts/times.ttf",
        ]
        for candidate in candidates:
            path = Path(candidate)
            if path.exists():
                pdfmetrics.registerFont(TTFont(font_name, str(path)))
                return font_name

        raise DocumentGenerationError(
            "Не найден TTF-шрифт с поддержкой кириллицы для генерации PDF. "
            "Установите DejaVu Sans, Arial или Times New Roman."
        )
