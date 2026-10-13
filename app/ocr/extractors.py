from __future__ import annotations

import asyncio
import tempfile
from dataclasses import dataclass
from pathlib import Path

from app.config import Settings
from app.ocr.tesseract import TesseractOCR


class DocumentExtractError(RuntimeError):
    pass


@dataclass(slots=True)
class DocumentExtractResult:
    text: str
    small_print: str | None


class DocumentTextExtractor:
    def __init__(self, settings: Settings, ocr: TesseractOCR) -> None:
        self._settings = settings
        self._ocr = ocr

    async def extract(self, file_path: Path, file_name: str) -> DocumentExtractResult:
        return await asyncio.to_thread(self._extract_sync, file_path, file_name)

    def _extract_sync(self, file_path: Path, file_name: str) -> DocumentExtractResult:
        suffix = Path(file_name).suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf(file_path)
        if suffix == ".docx":
            return self._extract_docx(file_path)
        if suffix in {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp"}:
            result = self._ocr.recognize_image_sync(file_path)
            return DocumentExtractResult(text=result.text, small_print=result.small_text)
        raise DocumentExtractError("Поддерживаются только PDF, DOCX и изображения.")

    def _extract_pdf(self, file_path: Path) -> DocumentExtractResult:
        import pdfplumber

        text_parts: list[str] = []
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        text_parts.append(page_text.strip())
        except Exception as exc:
            raise DocumentExtractError(f"Не удалось прочитать PDF: {exc}") from exc

        small_print = self._extract_pdf_small_print(file_path)
        text = "\n\n".join(text_parts).strip()
        if not text:
            # Сканированный PDF не имеет текстового слоя, поэтому рендерим страницы и запускаем OCR.
            ocr_result = self._ocr_pdf_pages(file_path)
            text = ocr_result.text
            small_print = small_print or ocr_result.small_print

        if not text:
            raise DocumentExtractError("Не удалось извлечь текст из PDF.")
        return DocumentExtractResult(text=text, small_print=small_print)

    def _extract_pdf_small_print(self, file_path: Path) -> str | None:
        import fitz

        fragments: list[str] = []
        try:
            document = fitz.open(file_path)
            try:
                for page in document:
                    blocks = page.get_text("dict").get("blocks", [])
                    for block in blocks:
                        for line in block.get("lines", []):
                            for span in line.get("spans", []):
                                text = str(span.get("text", "")).strip()
                                size = float(span.get("size", 0))
                                if text and 0 < size <= self._settings.small_font_pt:
                                    fragments.append(text)
            finally:
                document.close()
        except Exception:
            return None

        combined = " ".join(fragments)
        return combined[:5000] if combined else None

    def _ocr_pdf_pages(self, file_path: Path) -> DocumentExtractResult:
        import fitz

        text_parts: list[str] = []
        small_parts: list[str] = []
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                document = fitz.open(file_path)
                try:
                    for page_number, page in enumerate(document, start=1):
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        image_path = Path(temp_dir) / f"page_{page_number}.png"
                        pixmap.save(image_path)
                        result = self._ocr.recognize_image_sync(image_path)
                        text_parts.append(result.text)
                        if result.small_text:
                            small_parts.append(result.small_text)
                finally:
                    document.close()
        except Exception as exc:
            raise DocumentExtractError(f"Не удалось распознать сканированный PDF: {exc}") from exc

        return DocumentExtractResult(
            text="\n\n".join(text_parts).strip(),
            small_print="\n".join(small_parts).strip() or None,
        )

    def _extract_docx(self, file_path: Path) -> DocumentExtractResult:
        from docx import Document

        try:
            document = Document(file_path)
        except Exception as exc:
            raise DocumentExtractError(f"Не удалось прочитать DOCX: {exc}") from exc

        paragraphs: list[str] = []
        small_fragments: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                paragraphs.append(text)
            for run in paragraph.runs:
                run_text = run.text.strip()
                size = run.font.size
                if run_text and size is not None and size.pt <= self._settings.small_font_pt:
                    small_fragments.append(run_text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        text = "\n\n".join(paragraphs).strip()
        if not text:
            raise DocumentExtractError("DOCX не содержит распознаваемого текста.")
        small_print = " ".join(small_fragments)[:5000] if small_fragments else None
        return DocumentExtractResult(text=text, small_print=small_print)
